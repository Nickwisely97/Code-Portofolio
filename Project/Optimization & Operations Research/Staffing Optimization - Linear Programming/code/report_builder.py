"""
report_builder.py
Executive PowerPoint report for the monthly workforce roster, built on the
portfolio's shared Executive_Report_Template design system.

Usage from the notebook:
    from report_builder import build_executive_report
    path = build_executive_report(results, figures, output_dir="../result")
"""

import os
import sys

_TEMPLATE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "Executive_Report_Template")
sys.path.insert(0, _TEMPLATE_DIR)
from report_template import (
    NAVY, GOLD, BG_ALT, BODY_CLR, MARGIN,
    new_presentation, new_slide as _template_new_slide, add_text, add_stat_card, add_eyebrow,
    add_panel, add_table, add_footnote, save_report,
)
from pptx.util import Inches

KICKER = "EXECUTIVE REPORT   |   MONTHLY WORKFORCE ROSTER"
FOOTER = "Mixed-Integer Programming (PuLP / CBC)  |  January 2027"
TOTAL_PAGES = 7
SIDE_LEFT, SIDE_W = 6.28, 3.22


def new_slide(prs, title, page):
    return _template_new_slide(prs, KICKER, FOOTER, title, page, TOTAL_PAGES)


def _bullets(slide, left, top, width, height, lines, size=14):
    add_text(slide, left, top, width, height, [(line, size, False, BODY_CLR, False) for line in lines])


def build_executive_report(results, figures, output_dir):
    """
    results : dict -- agents, shifts, ot_shifts, covered_pct, short_hours, surplus_hours,
              required_hours, monthly_pay, overtime_pay, compliance (DataFrame),
              summary (agent_summary DataFrame), frontier (DataFrame),
              impact (rule_impact DataFrame), surplus_by_hour (Series: clock hour -> mean surplus)
    figures : dict of chart paths -- hourly_demand, roster, coverage_heatmap, agent_balance
    """
    r = results
    s = r["summary"]
    passed = int(r["compliance"]["OK"].sum())
    no_weekend = int(r["impact"].loc["Without the weekend-off rule", "agents"])
    top_surplus = r["surplus_by_hour"].sort_values(ascending=False).head(4)
    surplus_hours_text = ", ".join(f"{h:02d}:00" for h in sorted(top_surplus.index))
    n_rules = len(r["compliance"])
    prs = new_presentation()

    # --- 1. Executive summary
    s1 = new_slide(prs, "Executive Summary", 1)
    cards = [
        (f"{r['agents']}", "Agents needed"),
        (f"{r['covered_pct']:.1%}", "Demand hours covered"),
        (f"{passed}/{n_rules}", "Hard rules met"),
        (f"Rp {r['monthly_pay'] / 1e6:,.0f}M", "Pay this month"),
    ]
    for i, (value, label) in enumerate(cards):
        add_stat_card(s1, MARGIN + i * 2.29, 1.75, value, label)
    add_eyebrow(s1, MARGIN, 3.44, 4.39, "KEY FINDING")
    add_panel(s1, MARGIN, 3.78, 4.39, 2.64, accent=NAVY)
    _bullets(s1, MARGIN + 0.31, 3.95, 3.9, 2.4, [
        f"{r['agents']} agents cover January's demand with {r['shifts']} shifts, "
        f"{r['ot_shifts']} of them with overtime.",
        f"Only {r['short_hours']} of {r['required_hours']:,} required agent-hours is short, "
        "while every labor-law and company rule holds.",
    ], size=13)
    add_eyebrow(s1, MARGIN + 4.61, 3.44, 4.39, "WHAT DRIVES THE NUMBER")
    add_panel(s1, MARGIN + 4.61, 3.78, 4.39, 2.64, accent=GOLD)
    _bullets(s1, MARGIN + 4.92, 3.95, 3.9, 2.4, [
        f"The weekend-off rule: nobody may work both Saturday and Sunday, so each weekend needs "
        f"two separate crews. Without it, {no_weekend} agents would be enough.",
        "The model minimizes penalties (missing demand first); pay is reported, not optimized.",
    ], size=13)

    # --- 2. Rules
    s2 = new_slide(prs, "Hard Rules, Re-Checked on the Final Roster", 2)
    rows = [(row["Hard rule"], f"{row['Result']}", f"{row['Limit']}", "Yes" if row["OK"] else "NO")
            for _, row in r["compliance"].iterrows()]
    add_table(s2, MARGIN, 1.6, 9.0, ["Rule", "Worst case", "Limit", "Met"], rows,
              header_h=0.4, row_h=0.43, first_col_frac=0.52)

    # --- 3. Demand and shifts
    s3 = new_slide(prs, "Demand and the Three Shifts", 3)
    add_eyebrow(s3, MARGIN, 1.58, 9, "AGENTS NEEDED PER HOUR, SAME EVERY DAY")
    s3.shapes.add_picture(figures["hourly_demand"], Inches(MARGIN), Inches(1.92), width=Inches(5.9))
    add_eyebrow(s3, SIDE_LEFT + 0.3, 1.58, 2.9, "WHAT TO NOTICE")
    add_panel(s3, SIDE_LEFT + 0.3, 1.92, 2.92, 4.2, accent=GOLD)
    _bullets(s3, SIDE_LEFT + 0.55, 2.05, 2.55, 4.0, [
        "Demand peaks sharply at lunch (12:00) and dinner (18:00).",
        "Fixed 9-hour shifts cannot follow such sharp peaks without idle hours around them.",
        "No base shift covers 06:00-07:00; only overtime can.",
        "Demand is an assumption; a forecast would replace it.",
    ], size=12)

    # --- 4. Roster
    s4 = new_slide(prs, "The January Roster", 4)
    add_eyebrow(s4, MARGIN, 1.58, 9, "ONE ROW PER AGENT, ONE COLUMN PER DAY")
    s4.shapes.add_picture(figures["roster"], Inches(MARGIN), Inches(1.92), width=Inches(9.0))
    add_footnote(s4, MARGIN, 6.35, 9.0, "Full roster as a pivot table (agent x date, shift codes such as M11 or E18_A2) in result/roster_january_2027.csv.")

    # --- 5. Coverage
    s5 = new_slide(prs, "Coverage Hour by Hour", 5)
    add_eyebrow(s5, MARGIN, 1.58, 9, "ON DUTY MINUS REQUIRED, EVERY HOUR OF THE MONTH")
    s5.shapes.add_picture(figures["coverage_heatmap"], Inches(MARGIN), Inches(1.92), width=Inches(9.0))
    add_footnote(s5, MARGIN, 6.0, 9.0,
                 f"Short: {r['short_hours']} agent-hour(s); surplus: {r['surplus_hours']:,} agent-hours over the month. "
                 f"Surplus concentrates at {surplus_hours_text}, where shifts overlap.")

    # --- 6. Fairness
    s6 = new_slide(prs, "Fairness: Shift Mix and Overtime", 6)
    add_eyebrow(s6, MARGIN, 1.58, 5.6, "SHIFTS AND OVERTIME PER AGENT")
    s6.shapes.add_picture(figures["agent_balance"], Inches(MARGIN), Inches(1.92), height=Inches(4.3))
    add_eyebrow(s6, SIDE_LEFT, 1.58, SIDE_W, "SPREAD ACROSS AGENTS")
    add_stat_card(s6, SIDE_LEFT, 1.92, f"{s['days'].min()}-{s['days'].max()}", "Days worked per agent",
                  width=SIDE_W, accent=NAVY)
    add_stat_card(s6, SIDE_LEFT, 3.53, f"{s['overtime_hours'].min()}-{s['overtime_hours'].max()} h",
                  "Overtime per agent this month", width=SIDE_W, accent=GOLD)

    # --- 7. Recommendations
    s7 = new_slide(prs, "Recommendations and Next Steps", 7)
    add_eyebrow(s7, MARGIN, 1.58, 4.39, "RECOMMENDATIONS")
    add_panel(s7, MARGIN, 1.92, 4.39, 4.3, accent=NAVY)
    _bullets(s7, MARGIN + 0.31, 2.08, 3.9, 4.0, [
        f"Staff the store with {r['agents']} agents on M / E / N shifts.",
        f"Review the weekend-off rule: it alone adds {r['agents'] - no_weekend} agents. "
        "A rotating rule (e.g. one full weekend off per month) is worth testing with this model.",
        f"Stagger shift start times: idle hours pile up at {surplus_hours_text}.",
    ], size=13)
    add_eyebrow(s7, MARGIN + 4.61, 1.58, 4.39, "NEXT STEPS")
    add_panel(s7, MARGIN + 4.61, 1.92, 4.39, 4.3, accent=GOLD)
    _bullets(s7, MARGIN + 4.92, 2.08, 3.9, 4.0, [
        "Replace assumed demand with a forecast of customers per hour.",
        "Derive agents per hour from the queueing simulation (Self-Order Terminal project).",
        "Add leave requests, public holidays and agent preferences as extra penalties.",
    ], size=13)

    return save_report(prs, output_dir, "Executive_Workforce_Roster_Report")
